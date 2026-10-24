# -*- coding: utf-8 -*-
import csv, os, sys
sys.path.insert(0, r'D:\560\ITR-8 mobility test sender1\ITR-8 mobility test sender1\qc_analysis\pylibs')
import matplotlib
matplotlib.use('Agg')
from matplotlib import pyplot as plt
from matplotlib import font_manager
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.utils import get_column_letter

OUT = r'D:\560\ITR-8 mobility test sender1\ITR-8 mobility test sender1\qc_analysis\phy_rate_out'
CHART_DIR = os.path.join(OUT, 'charts')
os.makedirs(CHART_DIR, exist_ok=True)

for f in [r'C:/Windows/Fonts/msyh.ttc', r'C:/Windows/Fonts/simhei.ttf']:
    if os.path.exists(f):
        font_manager.fontManager.addfont(f)
        plt.rcParams['font.sans-serif'] = ['Microsoft YaHei', 'SimHei']
        break
plt.rcParams['axes.unicode_minus'] = False

TITLES = {
    'dl_rate': 'DL 速率 (PHY)', 'ul_rate': 'UL 速率 (PHY)',
    'dl_rb': 'DL RB 占用', 'ul_rb': 'UL RB 占用',
    'dl_mcs': 'DL MCS', 'ul_mcs': 'UL MCS',
    'dl_bler': 'DL BLER / 重传', 'ul_harq': 'UL HARQ',
    'cqi_ri': 'CQI / RI', 'snr': 'SNR (dB)', 'dci': 'DCI 次数', 'dl_mod': 'DL 调制分布',
}
SERIES_LABEL = {
    'dl_rate': {'tb_bytes':'调度TB字节','pass_bytes':'CRC通过字节','fail_bytes':'CRC失败字节','rate_mbps':'调度速率Mbps','pass_rate_mbps':'净速率Mbps'},
    'ul_rate': {'tb_bytes':'TB字节','rate_mbps':'速率Mbps'},
    'dl_rb':   {'rb_avg':'平均RB','rb_max':'最大RB','rb_usage_pct':'占用率%'},
    'ul_rb':   {'rb_avg':'平均RB','rb_max':'最大RB','rb_usage_pct':'占用率%'},
    'dl_mcs':  {'mcs_avg':'平均','mcs_min':'最小','mcs_max':'最大'},
    'ul_mcs':  {'mcs_avg':'平均','mcs_min':'最小','mcs_max':'最大'},
    'dl_bler': {'bler_pct':'BLER%','retx_pct':'重传%','crc_pass':'CRC通过','crc_fail':'CRC失败'},
    'ul_harq': {'new_tx':'新传','retx':'重传','retx_pct':'重传率%'},
    'cqi_ri':  {'cqi_avg':'CQI','ri_avg':'RI'},
    'snr':     {'snr_avg':'平均','rx0':'RX0','rx1':'RX1','rx2':'RX2','rx3':'RX3'},
    'dci':     {'ul_dci':'UL DCI','dl_dci':'DL DCI'},
    'dl_mod':  {'qpsk':'QPSK','qam16':'16QAM','qam64':'64QAM','qam256':'256QAM'},
}

def to_num(v):
    if v is None or str(v).strip() == '':
        return None
    try:
        return int(v)
    except ValueError:
        pass
    try:
        return float(v)
    except ValueError:
        return v  # 保留非数字文本

def read_csv(name):
    path = os.path.join(OUT, name + '.csv')
    header, rows = [], []
    with open(path, encoding='utf-8', newline='') as fh:
        rd = csv.reader(fh)
        header = next(rd)
        for r in rd:
            rows.append(r)
    return path, header, rows

wb = Workbook()
wb.remove(wb.active)

for name, title in TITLES.items():
    path, header, rows = read_csv(name)
    if not rows:
        continue
    val_cols = [i for i, h in enumerate(header) if h not in ('sec', 'time')]
    if not val_cols:
        continue
    times = [r[1] if len(r) > 1 else '' for r in rows]

    # ---------- PNG：每个数值列一个子图，各自量程 ----------
    n = len(val_cols)
    cols_wide = 2 if n > 1 else 1
    rows_n = (n + cols_wide - 1) // cols_wide
    fig, axes = plt.subplots(rows_n, cols_wide, figsize=(7*cols_wide, 3.2*rows_n), dpi=110, squeeze=False)
    for ax in axes.flat:
        ax.axis('off')
    for k, ci in enumerate(val_cols):
        ax = axes.flat[k]
        ax.axis('on')
        col = header[ci]
        ys = [to_num(r[ci]) for r in rows]
        xi, yi = [], []
        for i, y in enumerate(ys):
            if isinstance(y, (int, float)):
                xi.append(i)
                yi.append(y)
        ax.plot(xi, yi, marker='o', ms=3, lw=1.6, color='#1f77b4')
        ax.set_title(SERIES_LABEL.get(name, {}).get(col, col), fontsize=10)
        ax.grid(True, ls='--', alpha=0.5)
        if len(xi) > 8:
            step = max(1, len(xi)//6)
            ax.set_xticks(xi[::step])
            ax.set_xticklabels([times[i] for i in xi[::step]], rotation=45, fontsize=7)
        else:
            ax.set_xticks(xi)
            ax.set_xticklabels([times[i] for i in xi], rotation=45, fontsize=7)
    fig.suptitle(title, fontsize=12)
    fig.tight_layout(rect=[0, 0, 1, 0.95])
    fig.savefig(os.path.join(CHART_DIR, name + '.png'))
    plt.close(fig)

    # ---------- XLSX：数值类型化 + 每个数值列一张独立折线图 ----------
    ws = wb.create_sheet(title=name[:31])
    ws.append(header)
    for r in rows:
        ws.append([to_num(v) for v in r])
    anchor_col = get_column_letter(len(header) + 2)
    for k, ci in enumerate(val_cols):
        col = header[ci]
        chart = LineChart()
        chart.title = SERIES_LABEL.get(name, {}).get(col, col)
        chart.style = 2
        chart.height = 7
        chart.width = 18
        chart.y_axis.scaling.min = None
        chart.y_axis.scaling.max = None
        data = Reference(ws, min_col=ci+1, min_row=1, max_row=len(rows)+1)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(Reference(ws, min_col=2, min_row=2, max_row=len(rows)+1))
        ws.add_chart(chart, f'{anchor_col}{2 + k*17}')

xlsx = os.path.join(OUT, 'phy_rate_charts.xlsx')
wb.save(xlsx)
print('sheets:', len(wb.sheetnames))
print('xlsx :', xlsx)
print('pngs :', len(os.listdir(CHART_DIR)))
