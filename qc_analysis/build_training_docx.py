# -*- coding: utf-8 -*-
import os, re
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import parse_xml

DOC_DIR = r'D:\560\ITR-8 mobility test sender1\ITR-8 mobility test sender1\培训文档'
DOCX = os.path.join(DOC_DIR, 'Perl工具解析QXDM日志与AI平台协议分析培训手册.docx')

doc = Document()
doc.core_properties.title = 'Perl工具解析QXDM日志与AI平台协议分析培训手册'
doc.core_properties.author = ''

# 默认字体
normal = doc.styles['Normal']
normal.font.name = 'Microsoft YaHei'
normal.font.size = Pt(10.5)
normal._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')

def set_east(run, name='Microsoft YaHei'):
    run.font.name = name
    run._element.rPr.rFonts.set(qn('w:ascii'), name)
    run._element.rPr.rFonts.set(qn('w:hAnsi'), name)
    run._element.rPr.rFonts.set(qn('w:eastAsia'), name)

def para(text, bold=False, size=10.5, style=None):
    p = doc.add_paragraph(style=style)
    r = p.add_run(text)
    r.bold = bold
    r.font.size = Pt(size)
    set_east(r)
    return p

def h1(text):
    p = doc.add_heading(level=1)
    r = p.add_run(text)
    r.font.size = Pt(15)
    set_east(r)
    r.font.color.rgb = RGBColor(0x1F, 0x3B, 0x64)
    return p

def h2(text):
    p = doc.add_heading(level=2)
    r = p.add_run(text)
    r.font.size = Pt(12)
    set_east(r)
    r.font.color.rgb = RGBColor(0x2E, 0x54, 0x86)
    return p

def bullets(items):
    for it in items:
        p = doc.add_paragraph(style='List Bullet')
        r = p.add_run(it)
        r.font.size = Pt(10.5)
        set_east(r)

def code_block(text):
    t = doc.add_table(rows=1, cols=1)
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = t.cell(0, 0)
    cell._tc.get_or_add_tcPr().append(
        parse_xml(r'<w:shd %s w:val="clear" w:color="auto" w:fill="F5F5F5"/>' % ('xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"')))
    lines = text.split('\n')
    first = True
    for ln in lines:
        p = cell.paragraphs[0] if first else cell.add_paragraph()
        first = False
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_before = Pt(0)
        r = p.add_run(ln if ln else ' ')
        r.font.name = 'Consolas'
        r.font.size = Pt(8)
        r._element.rPr.rFonts.set(qn('w:ascii'), 'Consolas')
        r._element.rPr.rFonts.set(qn('w:hAnsi'), 'Consolas')
        r._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    return t

def msg_table(rows):
    t = doc.add_table(rows=1, cols=3)
    t.style = 'Table Grid'
    hdr = t.rows[0].cells
    for i, htxt in enumerate(['分组', '消息 (0x代码)', '作用']):
        p = hdr[i].paragraphs[0]
        r = p.add_run(htxt)
        r.bold = True
        r.font.size = Pt(9)
        set_east(r)
    for row in rows:
        cells = t.add_row().cells
        for i, txt in enumerate(row):
            p = cells[i].paragraphs[0]
            r = p.add_run(txt)
            r.font.size = Pt(9)
            if i == 1:
                r.font.name = 'Consolas'
                r._element.rPr.rFonts.set(qn('w:ascii'), 'Consolas')
                r._element.rPr.rFonts.set(qn('w:hAnsi'), 'Consolas')
            else:
                set_east(r)
    return t

# ---------------- 封面标题 ----------------
tp = doc.add_heading(level=0)
tr = tp.add_run('Perl工具解析QXDM日志与AI平台协议分析培训手册')
tr.font.size = Pt(22)
set_east(tr)
tr.font.color.rgb = RGBColor(0, 0, 0)

sub = doc.add_paragraph()
sr = sub.add_run('日志过滤篇：速率 / 接入 / 掉线三类问题的日志提取')
sr.font.size = Pt(12)
sr.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
set_east(sr)

para('本手册面向需要使用 Perl 脚本从 QXDM 抓取的 NR5G 日志中快速提取特定问题域报文、再把提取结果交给 AI 平台做协议分析的工程师。'
     '手册讲解三个过滤脚本：速率、接入、掉线，每个脚本只做消息过滤，不做时间过滤，整个日志范围导出，'
     '输出为带时间戳的纯文本解码流，可直接粘贴给 AI 平台追问协议细节。三个脚本的 .pl 文件与本手册放在同一目录。')

# ---------------- 1 培训目的 ----------------
h1('1 培训目的与适用范围')
para('QXDM 原始日志报文量极大（动辄上千万条），直接交给 AI 平台既慢又容易被无关消息淹没。'
     '先用 Perl 调用 QCAT 的消息过滤器，把与问题域相关的消息类型挑出来、按时间顺序导出成文本，'
     '再交给 AI 平台，AI 才能集中精力分析协议行为。本手册覆盖三类最常见的问题：上下行速率、接入（搜网注册附着）、掉线（无线链路失败）。')
bullets([
    '能独立运行三个 Perl 过滤脚本并看懂输出文件。',
    '知道每类问题需要哪些 NR5G 消息类型、每条消息的作用。',
    '能把过滤后的文本连同问题描述一起交给 AI 平台，获得协议级分析。',
])

# ---------------- 2 总体流程 ----------------
h1('2 工作流程与总体思路')
bullets([
    '第一步：用 QXDM 抓取日志并保存为 QCAT 可打开的格式（HDF / QMDL / DLF）。',
    '第二步：修改脚本顶部配置区的输入输出路径（本手册脚本不含时间过滤，导出整个日志）。',
    '第三步：在装有 QCAT 6.x 的电脑上双击或命令行运行 perl 脚本，等待生成 .txt 文件。',
    '第四步：按问题时间点截取 .txt 中的日志块（块头格式为 === 序号 0x消息类型 ===），连同问题现象交给 AI 平台。',
])
para('一条重要经验：设备日志时间为 UTC，与本地时间可能相差数小时（如北京时间差 8 小时），'
     '向 AI 平台描述问题时务必说明日志时区，否则会对应错事件。', bold=True)

# ---------------- 3 环境准备 ----------------
h1('3 环境准备')
h2('3.1 软件要求')
bullets([
    'QCAT 6.x（安装后提供 COM 组件，脚本通过 QCAT6.Application 调用）。',
    'Strawberry Perl（含 Win32::OLE 模块），脚本用 perl 解释器直接运行。',
])
h2('3.2 运行方法')
code_block('''# 直接运行（使用脚本内配置区路径）
perl 01_filter_rate.pl

# 输出示例（过滤后的文本）
=== 00000001  0xB823 ===
2026 Oct 22  01:17:51.259  [BE]  0xB823  NR5G RRC Serving Cell Info
Physical Cell ID = 404
DL Frequency = 627264
DL Bandwidth = 100
...
# ============ 速率日志摘要 ============
# 按消息类型统计
# 0xB888  23547  NR5G MAC PDSCH Stats''')
h2('3.3 输出文件格式说明')
bullets([
    '每条消息前有块头 === 八位序号 0x类型 ===，序号即消息在过滤结果中的顺序。',
    '块头下一行是原始时间戳（设备时间），之后是消息的完整解码文本。',
    '文件末尾自动附加事件摘要与按消息类型的统计，便于快速判断日志里有什么。',
])
para('本手册三个脚本均不包含时间过滤，目的是拿到完整证据链；若后续需要按时间窗提取，'
     '可在脚本中增加 QCAT 的 SetTimeWindow 调用（不在本手册展开）。')

# ---------------- 4 速率 ----------------
h1('4 工具一 过滤分析5G上下行速率日志')
para('用途：提取上下行调度、RB 占用、MCS、层数、调制方式、BLER 重传、信号质量与数据面字节数，'
     '用于定位速率低、速率波动、RB 资源不足、BLER 高等问题。')
msg_table([
    ('调度', '0xB885 / 0xB883 / 0xB886 / 0xB887', 'DCI 授权、PUSCH/PDSCH 逐次调度的 RB/MCS/TB Size'),
    ('CRC 统计', '0xB888', '下行 PDSCH CRC 通过/失败累计、BLER、重传'),
    ('反馈', '0xB89B / 0xB896 / 0xB8A7', 'HARQ ACK/NACK、CQI/RI/PMI'),
    ('射频', '0xB8D8 / 0xB8B5', 'Serving SNR、上行发射功率分布'),
    ('数据面', '0xB857 / 0xB872 / 0xB842 / 0xB860', 'L2 上下行 TB 字节、PDCP 上下行字节'),
    ('小区信息', '0xB823', '服务小区 PCI、频点、带宽'),
])
h2('4.1 输出判读要点')
bullets([
    '下行物理层速率用 0xB888 的 CRC Pass TB Bytes 累计值差分计算，或用 0xB887 每条 PDSCH 的 TB Size 求和。',
    '上行速率用 0xB883 中 PUSCH 记录的 TB Size 求和。',
    'RB 占用、MCS、调制、层数看 0xB887（下行）与 0xB883（上行）的对应字段，判断调度是否喂饱。',
    'BLER 持续高、MCS 被压低，说明链路质量差，结合 0xB8D8 的 SNR 一起看。',
])

# ---------------- 5 接入 ----------------
h1('5 工具二 过滤分析5G接入问题日志')
para('用途：提取 RRC 建链与 NAS 注册流程、搜网与小区评估、随机接入全过程，'
     '用于分析搜网慢、无服务、Attach/Registration 被拒绝、注册后不建业务等问题。')
msg_table([
    ('RRC 信令', '0xB821 / 0xB822 / 0xB823 / 0xB825', 'RRC Setup/重配/拒绝、MIB、服务小区、RRC 配置'),
    ('NAS 信令', '0xB80C / 0xB80D', 'MM5G 注册状态、Service Request'),
    ('随机接入', '0xB889 / 0xB88A', 'RACH 触发原因与 Msg1~Msg4 结果'),
    ('搜网评估', '0xB970 / 0xB96F / 0xB950 / 0xB96D / 0xB98F / 0xB9A7', '空闲态 S 准则、连接态评估、公共配置、天线分集'),
    ('辅助确认', '0xB890 / 0xB842 / 0xB860 / 0xB8D8', 'DRX、SRB 信令流量、SNR'),
])
h2('5.1 输出判读要点')
bullets([
    '在 RRC OTA 文本中检索 Attach reject / Registration reject，看 EMM/5GMM cause（如 cause 27 N1 mode not allowed 表示 N1 被禁）。',
    '0xB80C 看注册状态是否进入注册态；0xB889/0xB88A 看随机接入是否成功、失败在哪一步（Msg1/Msg2/Msg3/Msg4）。',
    '连续 RACH 失败且 0xB823 服务小区不变化时，优先怀疑上行链路或小区阻塞。',
])

# ---------------- 6 掉线 ----------------
h1('6 工具三 过滤分析5G掉线问题日志')
para('用途：提取无线链路失败（RLF）前后的一条完整证据链：RRC 与随机接入、波束失败与 RLM、'
     '上行发送挂起、信号质量、网络侧调度、L2 缓冲，用于定位语音/数据掉线的根因。')
msg_table([
    ('RRC/RLF', '0xB821 / 0xB822 / 0xB823 / 0xB825 / 0xB889 / 0xB88A', '重配置、服务小区变化、RACH 触发原因（RLF/BEAM_FAILURE）'),
    ('波束与 RLM', '0xB959 / 0xB8AE / 0xB89C / 0xB890', 'IS/OOS BLER、TX_SUSPEND 上行丢弃、流控、DRX'),
    ('信号质量', '0xB8D8 / 0xB97F / 0xB8B5', 'SNR、RSRP/RSRQ/SINR、发射功率'),
    ('调度反馈', '0xB885 / 0xB883 / 0xB888 / 0xB89B / 0xB896', 'DCI 是否还在下发、PDSCH 统计、HARQ 反馈'),
    ('L2 缓冲', '0xB873 / 0xB860 / 0xB842 / 0xB868 / 0xB84D / 0xB872 / 0xB857', 'BSR、PDCP/RLC 收发与缓冲堆积'),
])
h2('6.1 输出判读要点')
bullets([
    '先看 0xB889 RACH Trigger 的原因字段，区分 RLF、波束失败、SR 失败。',
    '0xB959 的 Purpose=BEAM_FAILURE 与 IS/OOS BLER 记录波束失败检测过程。',
    '0xB8AE 的 Drop Reason=TX_SUSPEND 表示 MAC 丢弃 PUCCH/PUSCH 发送，是上行被挂起的直接证据。',
    '对照 0xB885 的 DCI 次数：网络侧不再下发授权时，BSR/SR 发不出去，缓冲在 0xB860/0xB868 持续堆积。',
])

# ---------------- 7 AI 平台 ----------------
h1('7 输出结果交给AI平台做协议分析')
h2('7.1 提问模板')
para('过滤后的 .txt 是纯文本，直接粘贴或上传给 AI 平台即可。建议按下面模板提问：')
code_block('''【日志背景】NR5G，设备日志为 UTC 时间；小区 PCI=347，N78，100MHz。
【问题现象】HH:MM:SS 至 HH:MM:SS 期间语音中断，约 20 秒后 RRC 重建恢复。
【过滤日志】已过滤 RACH Trigger/Attempt、Skip UL TX、RLM、DCI、PDCP/RLC 统计。
【问题】
1. 该时段 MAC 是否处于 SR/BSR 待处理状态？为什么 PUCCH/PUSCH 未发送？
2. TX_SUSPEND 由什么条件触发、如何清除？
3. 网络侧最后一条上行授权与最后一次 SR 分别在什么时间？''')
h2('7.2 三个场景的 AI 分析指引文件（SKILL.md）')
para('手册同目录另附三个 SKILL.md 文件。把对应场景的 SKILL.md 内容与过滤后的日志一起提供给 AI 平台，'
     'AI 会按其中的消息速查与判定规则，输出带日志块号证据的分析报告。')
bullets([
    '01_throughput_analysis_skill.md：上下行速率分析。',
    '02_access_analysis_skill.md：接入（搜网/注册/RACH）分析。',
    '03_drop_analysis_skill.md：掉线/RLF 分析。',
])
bullets([
    '提问时把时间范围、PCI、频点写在背景里，AI 才能对齐事件。',
    '日志块保留 === 序号 0x类型 === 块头，AI 引用块号时你能快速回查。',
    '让 AI 先给事件时间线，再给根因结论，并要求每个结论都引用日志块号。',
])

# ---------------- 附录 ----------------
for app, path, name in [
    ('A', '01_filter_rate.pl', '过滤分析5G上下行速率日志'),
    ('B', '02_filter_access.pl', '过滤分析5G接入问题日志'),
    ('C', '03_filter_drop.pl', '过滤分析5G掉线问题日志'),
]:
    h1(f'附录{app} {name} 完整源码 ({path})')
    para('脚本与本文档同目录存放，可直接用文本编辑器另存为同名 .pl 文件；运行前修改配置区的输入输出路径。')
    src = open(os.path.join(DOC_DIR, path), encoding='utf-8').read()
    code_block(src)

for app, path, name in [
    ('D', '01_throughput_analysis_skill.md', '5G上下行速率场景 AI 分析指引'),
    ('E', '02_access_analysis_skill.md', '5G接入场景 AI 分析指引'),
    ('F', '03_drop_analysis_skill.md', '5G掉线场景 AI 分析指引'),
]:
    h1(f'附录{app} {name} ({path})')
    para('将本文件内容与过滤后的日志一起提供给 AI 平台作为分析指引，AI 会按其中的工作流、消息速查与判定规则输出结论。')
    src = open(os.path.join(DOC_DIR, path), encoding='utf-8').read()
    code_block(src)

doc.save(DOCX)
print('saved', DOCX)
