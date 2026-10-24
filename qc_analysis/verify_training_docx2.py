# -*- coding: utf-8 -*-
import os
from docx import Document
DIR = r'D:\560\ITR-8 mobility test sender1\ITR-8 mobility test sender1\培训文档'
DOCX = os.path.join(DIR, 'Perl工具解析QXDM日志与AI平台协议分析培训手册.docx')
doc = Document(DOCX)
texts = []
for t in doc.tables:
    texts.append('\n'.join(c.text for c in t.cell(0,0).paragraphs))
def norm(s):
    return '\n'.join(l.rstrip().replace(' ', '', 1) if l.strip() == '' else l.rstrip() for l in s.replace('\r', '').split('\n')).strip()
def norm2(s):
    return '\n'.join(l.rstrip() for l in s.replace('\r', '').split('\n')).strip()
for name in ['01_filter_rate.pl', '02_filter_access.pl', '03_filter_drop.pl']:
    src = open(os.path.join(DIR, name), encoding='utf-8').read()
    hits = [i for i, t in enumerate(texts) if norm(t) == norm2(src)]
    if not hits:
        # 找最接近的表，打印差异行
        cand = None; best = 0
        for i, t in enumerate(texts):
            a, b = norm(t).split('\n'), norm2(src).split('\n')
            common = sum(1 for x, y in zip(a, b) if x == y)
            if common > best:
                best, cand = common, i
        a = norm(texts[cand]).split('\n'); b = norm2(src).split('\n')
        print(name, 'MISS, best common', best, '/', len(b), 'table', cand)
        for j, (x, y) in enumerate(zip(a, b)):
            if x != y:
                print('  line', j, 'DOCX:', repr(x)[:70], ' SRC:', repr(y)[:70])
                if j > 6: break
    else:
        print(name, 'embedded OK (table idx', hits, ')')
