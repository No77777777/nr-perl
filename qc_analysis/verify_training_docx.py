# -*- coding: utf-8 -*-
import os, re, zipfile
from docx import Document
DIR = r'D:\560\ITR-8 mobility test sender1\ITR-8 mobility test sender1\培训文档'
DOCX = os.path.join(DIR, 'Perl工具解析QXDM日志与AI平台协议分析培训手册.docx')
doc = Document(DOCX)
heads = [p.text for p in doc.paragraphs if p.style.name.startswith(('Title','Heading'))]
print('HEADINGS:')
for h in heads: print('  ', h)
print('paragraphs:', len(doc.paragraphs), ' tables:', len(doc.tables))

# 校验附录代码与源文件一致
texts = []
for t in doc.tables:
    txt = '\n'.join(c.text for c in t.cell(0,0).paragraphs)
    texts.append(txt)
for name in ['01_filter_rate.pl','02_filter_access.pl','03_filter_drop.pl']:
    src = open(os.path.join(DIR, name), encoding='utf-8').read()
    found = any(t.replace('\r','') == src.replace('\r','') for t in texts)
    print(name, 'embedded-matches-file:', found)
print('citation tokens in doc:', any(':codex' in t for t in texts))

# XML 合法性
z = zipfile.ZipFile(DOCX)
bad = z.testzip()
print('zip ok:', bad is None)
import xml.dom.minidom
for n in ['word/document.xml']:
    xml.dom.minidom.parseString(z.read(n))
    print(n, 'well-formed')
