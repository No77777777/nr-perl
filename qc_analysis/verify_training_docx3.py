# -*- coding: utf-8 -*-
import os, zipfile, xml.dom.minidom
from docx import Document
DIR = r'D:\560\ITR-8 mobility test sender1\ITR-8 mobility test sender1\培训文档'
DOCX = os.path.join(DIR, 'Perl工具解析QXDM日志与AI平台协议分析培训手册.docx')
doc = Document(DOCX)
heads = [p.text for p in doc.paragraphs if p.style.name.startswith(('Title','Heading'))]
print('HEADINGS:')
for h in heads: print('  ', h)
print('tables:', len(doc.tables), ' paragraphs:', len(doc.paragraphs))

def norm(s):
    return '\n'.join(l.rstrip() for l in s.replace('\r','').split('\n')).strip()
texts = ['\n'.join(c.text for c in t.cell(0,0).paragraphs) for t in doc.tables]
for name in ['01_throughput_analysis_skill.md','02_access_analysis_skill.md','03_drop_analysis_skill.md']:
    src = open(os.path.join(DIR, name), encoding='utf-8').read()
    print(name, 'embedded OK:', any(norm(t).replace(' ','') == norm(src).replace(' ','') for t in texts))
z = zipfile.ZipFile(DOCX)
print('zip ok:', z.testzip() is None)
xml.dom.minidom.parseString(z.read('word/document.xml'))
print('document.xml well-formed')
